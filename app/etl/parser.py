"""DOCX、文本型 PDF、Excel 及文本格式的安全解析器。"""

from __future__ import annotations

from io import BytesIO
import hashlib
import json
from pathlib import Path, PurePosixPath
import re
import shutil
import subprocess
import tempfile
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import xlrd

from app.etl.contracts import Chunk, ParsedBlock, ParsedDocument


MAX_ARCHIVE_ENTRIES = 5000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_PDF_PAGES = 100
MAX_OCR_PAGES = 30
MAX_SHEETS = 20
MAX_CELLS = 200000
MAX_CHARACTERS = 1000000
SUPPORTED = {".json", ".jsonl", ".txt", ".md", ".docx", ".pdf", ".xlsx", ".xls"}


class DocumentParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _clean(value: object) -> str:
    return re.sub(r"[ \t]+", " ", str(value or "").replace("\x00", "")).strip()


def _archive_guard(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise DocumentParseError("archive_too_large", "压缩包条目过多")
            total = 0
            for item in entries:
                path = PurePosixPath(item.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise DocumentParseError("malformed_document", "压缩包包含非法路径")
                total += item.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise DocumentParseError("archive_too_large", "解压后文件过大")
                if item.compress_size and item.file_size / item.compress_size > MAX_COMPRESSION_RATIO:
                    raise DocumentParseError("archive_too_large", "压缩比异常")
    except zipfile.BadZipFile as error:
        raise DocumentParseError("malformed_document", "Office 文件结构损坏") from error


def _text_blocks(text: str, kind: str = "paragraph") -> list[ParsedBlock]:
    return [ParsedBlock(kind=kind, text=line) for line in (_clean(part) for part in text.splitlines()) if line]


def _parse_json(data: bytes, jsonl: bool) -> list[ParsedBlock]:
    try:
        text = data.decode("utf-8-sig")
        values = [json.loads(line) for line in text.splitlines() if line.strip()] if jsonl else [json.loads(text)]
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise DocumentParseError("malformed_document", "JSON 编码或结构不合法") from error
    blocks = []
    for index, value in enumerate(values, 1):
        rendered = json.dumps(value, ensure_ascii=False, sort_keys=True)
        blocks.append(ParsedBlock(kind="record", text=rendered, location={"record": index}))
    return blocks


def _parse_docx(data: bytes) -> tuple[list[ParsedBlock], list[str]]:
    _archive_guard(data)
    try:
        document = Document(BytesIO(data))
    except Exception as error:
        raise DocumentParseError("malformed_document", "DOCX 无法读取") from error
    blocks: list[ParsedBlock] = []
    warnings: list[str] = []
    for order, item in enumerate(document.iter_inner_content(), 1):
        if hasattr(item, "rows"):
            rows = []
            for row in item.rows:
                values = [_clean(cell.text) for cell in row.cells]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                blocks.append(ParsedBlock(kind="table", text="\n".join(rows), location={"order": order}))
        else:
            text = _clean(item.text)
            if text:
                style = _clean(getattr(getattr(item, "style", None), "name", ""))
                kind = "heading" if style.lower().startswith("heading") or style.startswith("标题") else "paragraph"
                blocks.append(ParsedBlock(kind=kind, text=text, location={"order": order, "style": style}))
    if document.inline_shapes:
        warnings.append("文档中的图片未执行 OCR")
    return blocks, warnings


def _ocr_pdf(data: bytes, page_count: int) -> tuple[list[ParsedBlock], list[str]]:
    if page_count > MAX_OCR_PAGES:
        raise DocumentParseError("ocr_too_many_pages", f"扫描PDF OCR最多支持 {MAX_OCR_PAGES} 页")
    pdftoppm, tesseract = shutil.which("pdftoppm"), shutil.which("tesseract")
    if not pdftoppm or not tesseract:
        raise DocumentParseError("ocr_unavailable", "服务器未安装OCR运行依赖")
    blocks: list[ParsedBlock] = []
    with tempfile.TemporaryDirectory(prefix="bank-ocr-") as temporary:
        root = Path(temporary)
        source = root / "input.pdf"
        source.write_bytes(data)
        try:
            subprocess.run(
                [pdftoppm, "-png", "-r", "200", "-f", "1", "-l", str(page_count),
                 str(source), str(root / "page")],
                check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=120,
            )
        except (subprocess.SubprocessError, OSError) as error:
            raise DocumentParseError("ocr_conversion_failed", "扫描PDF页面转换失败") from error
        images = sorted(root.glob("page-*.png"), key=lambda path: int(re.search(r"(\d+)$", path.stem).group(1)))
        for page_number, image in enumerate(images, 1):
            try:
                result = subprocess.run(
                    [tesseract, str(image), "stdout", "-l", "chi_sim+eng", "--psm", "6"],
                    check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                    text=True, encoding="utf-8", errors="replace", timeout=45,
                )
            except (subprocess.SubprocessError, OSError) as error:
                raise DocumentParseError("ocr_recognition_failed", f"PDF第 {page_number} 页OCR失败") from error
            text = _clean(result.stdout)
            if text:
                blocks.append(ParsedBlock(kind="ocr", text=text, location={
                    "page": page_number, "ocr": True, "language": "chi_sim+eng",
                }))
    if not blocks:
        raise DocumentParseError("ocr_no_content", "OCR未识别到可用文字")
    return blocks, ["扫描PDF已执行OCR，发布前必须人工核对识别结果"]


def _parse_pdf(data: bytes, allow_ocr: bool = False) -> tuple[list[ParsedBlock], list[str]]:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
    except Exception as error:
        raise DocumentParseError("malformed_document", "PDF 无法读取") from error
    if reader.is_encrypted:
        raise DocumentParseError("password_protected", "加密 PDF 暂不支持")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentParseError("too_many_pages", f"PDF 超过 {MAX_PDF_PAGES} 页")
    blocks: list[ParsedBlock] = []
    empty_pages = 0
    for page_number, page in enumerate(reader.pages, 1):
        try:
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > 8 * 1024 * 1024:
                raise DocumentParseError("page_too_large", "PDF 单页内容流过大")
            text = _clean(page.extract_text(extraction_mode="layout") or "") if contents is not None else ""
        except DocumentParseError:
            raise
        except Exception as error:
            raise DocumentParseError("malformed_document", f"PDF 第 {page_number} 页解析失败") from error
        if text:
            blocks.append(ParsedBlock(kind="paragraph", text=text, location={"page": page_number}))
        else:
            empty_pages += 1
    if not blocks:
        if allow_ocr:
            return _ocr_pdf(data, len(reader.pages))
        raise DocumentParseError("ocr_required", "PDF 未检测到可提取文字，需要 OCR")
    warnings = [f"{empty_pages} 页未检测到文字"] if empty_pages else []
    return blocks, warnings


def _row_text(values: list[object]) -> str:
    return " | ".join(_clean(value) for value in values).rstrip(" |")


def _parse_xlsx(data: bytes) -> tuple[list[ParsedBlock], list[str]]:
    _archive_guard(data)
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as error:
        raise DocumentParseError("malformed_document", "XLSX 无法读取") from error
    blocks: list[ParsedBlock] = []
    warnings: list[str] = []
    cells = 0
    try:
        worksheets = workbook.worksheets
        if len(worksheets) > MAX_SHEETS:
            raise DocumentParseError("too_many_sheets", f"Excel 超过 {MAX_SHEETS} 个工作表")
        for sheet in worksheets:
            if sheet.sheet_state != "visible":
                warnings.append(f"隐藏工作表未导入：{sheet.title}")
                continue
            rows = []
            start = None
            end = None
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise DocumentParseError("too_many_cells", f"Excel 超过 {MAX_CELLS} 个单元格")
                rendered = _row_text(list(row))
                if rendered:
                    rows.append(rendered)
                    start = start or row_number
                    end = row_number
                if len(rows) >= 100:
                    blocks.append(ParsedBlock(kind="sheet", text="\n".join(rows), location={
                        "sheet": sheet.title, "row_start": start, "row_end": end,
                    }))
                    rows, start, end = [], None, None
            if rows:
                blocks.append(ParsedBlock(kind="sheet", text="\n".join(rows), location={
                    "sheet": sheet.title, "row_start": start, "row_end": end,
                }))
    finally:
        workbook.close()
    return blocks, warnings


def _parse_xls(data: bytes) -> tuple[list[ParsedBlock], list[str]]:
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as error:
        raise DocumentParseError("malformed_document", "XLS 无法读取或受密码保护") from error
    if workbook.nsheets > MAX_SHEETS:
        raise DocumentParseError("too_many_sheets", f"Excel 超过 {MAX_SHEETS} 个工作表")
    blocks: list[ParsedBlock] = []
    cells = 0
    for index in range(workbook.nsheets):
        sheet = workbook.sheet_by_index(index)
        rows = []
        start = None
        for row_number in range(sheet.nrows):
            cells += sheet.ncols
            if cells > MAX_CELLS:
                raise DocumentParseError("too_many_cells", f"Excel 超过 {MAX_CELLS} 个单元格")
            rendered = _row_text(sheet.row_values(row_number))
            if rendered:
                rows.append(rendered)
                start = start if start is not None else row_number + 1
            if len(rows) >= 100:
                blocks.append(ParsedBlock(kind="sheet", text="\n".join(rows), location={
                    "sheet": sheet.name, "row_start": start, "row_end": row_number + 1,
                }))
                rows, start = [], None
        if rows:
            blocks.append(ParsedBlock(kind="sheet", text="\n".join(rows), location={
                "sheet": sheet.name, "row_start": start, "row_end": sheet.nrows,
            }))
        workbook.unload_sheet(index)
    workbook.release_resources()
    return blocks, ["XLS 仅读取公式缓存值，不执行公式"]


def parse_document(filename: str, data: bytes, allow_ocr: bool = False) -> ParsedDocument:
    safe_name = Path(filename).name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in SUPPORTED:
        raise DocumentParseError("unsupported_format", f"不支持的文件格式：{suffix or '无扩展名'}")
    if not data:
        raise DocumentParseError("malformed_document", "文件为空")
    warnings: list[str] = []
    if suffix in {".txt", ".md"}:
        try:
            blocks = _text_blocks(data.decode("utf-8-sig"))
        except UnicodeDecodeError as error:
            raise DocumentParseError("malformed_document", "文本文件必须使用 UTF-8") from error
    elif suffix in {".json", ".jsonl"}:
        blocks = _parse_json(data, suffix == ".jsonl")
    elif suffix == ".docx":
        blocks, warnings = _parse_docx(data)
    elif suffix == ".pdf":
        blocks, warnings = _parse_pdf(data, allow_ocr=allow_ocr)
    elif suffix == ".xlsx":
        blocks, warnings = _parse_xlsx(data)
    else:
        blocks, warnings = _parse_xls(data)
    if not blocks:
        raise DocumentParseError("no_content", "未解析到可发布内容")
    total = sum(len(block.text) for block in blocks)
    if total > MAX_CHARACTERS:
        raise DocumentParseError("too_many_characters", "解析结果超过字符上限")
    # Office ZIP 元数据含创建时间，直接对原始字节取哈希会让同一内容偶发失去幂等性。
    canonical = json.dumps(
        {"source_type": suffix.lstrip("."),
         "blocks": [{"kind": block.kind, "text": block.text} for block in blocks]},
        ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")
    return ParsedDocument(
        filename=safe_name,
        source_type=suffix.lstrip("."),
        title=Path(safe_name).stem[:200],
        sha256=hashlib.sha256(canonical).hexdigest(),
        blocks=blocks,
        warnings=warnings,
    )


def chunk_document(document: ParsedDocument, maximum: int = 1200, overlap: int = 120) -> list[Chunk]:
    if maximum < 200 or overlap < 0 or overlap >= maximum:
        raise ValueError("分块参数不合法")
    chunks: list[Chunk] = []
    for block in document.blocks:
        text = block.text
        start = 0
        while start < len(text):
            end = min(len(text), start + maximum)
            if end < len(text):
                boundary = max(text.rfind("\n", start, end), text.rfind("。", start, end))
                if boundary > start + maximum // 2:
                    end = boundary + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append(Chunk(index=len(chunks), text=piece, source={
                    "filename": document.filename,
                    "source_type": document.source_type,
                    "sha256": document.sha256,
                    **block.location,
                }))
            if end >= len(text):
                break
            start = end - overlap
    return chunks
