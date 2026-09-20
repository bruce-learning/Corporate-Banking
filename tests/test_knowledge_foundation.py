from io import BytesIO
import json

import httpx
from docx import Document
from openpyxl import Workbook
from pydantic import SecretStr
from pypdf import PdfWriter
import pytest

from app.config import Settings
from app.etl.parser import DocumentParseError, chunk_document, parse_document
from app.infrastructure.llm.aliyun import AliyunModelClient, ModelServiceError, PaidCallsDisabled
from app.infrastructure.llm.contracts import ChatInput, ModelMessage
from scripts.generate_knowledge_dataset import generate, generate_rows
from tests.conftest import login


def docx_bytes() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("企业网银", level=1)
    document.add_paragraph("企业网银支持查询和转账。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "规则"
    table.cell(1, 0).text = "收费"
    table.cell(1, 1).text = "虚构演示"
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收费规则"
    sheet.append(["项目", "金额"])
    sheet.append(["演示服务费", 120])
    workbook.save(buffer)
    return buffer.getvalue()


def test_docx_xlsx_and_text_parsing():
    docx = parse_document("规则.docx", docx_bytes())
    assert docx.source_type == "docx"
    assert any(block.kind == "heading" for block in docx.blocks)
    assert any(block.kind == "table" and "虚构演示" in block.text for block in docx.blocks)
    assert chunk_document(docx)[0].source["filename"] == "规则.docx"

    xlsx = parse_document("收费.xlsx", xlsx_bytes())
    assert xlsx.source_type == "xlsx"
    assert xlsx.blocks[0].location["sheet"] == "收费规则"
    assert "演示服务费" in xlsx.blocks[0].text

    text = parse_document("说明.md", "# 标题\n这是合成说明。".encode())
    assert text.character_count > 0


def test_pdf_without_text_is_explicit_ocr_required():
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    with pytest.raises(DocumentParseError) as caught:
        parse_document("scan.pdf", buffer.getvalue())
    assert caught.value.code == "ocr_required"


def test_unsupported_and_malformed_files_are_rejected():
    with pytest.raises(DocumentParseError) as caught:
        parse_document("legacy.doc", b"not-a-document")
    assert caught.value.code == "unsupported_format"
    with pytest.raises(DocumentParseError) as caught:
        parse_document("broken.docx", b"not-a-zip")
    assert caught.value.code == "malformed_document"


def test_dataset_counts_splits_and_manifest(tmp_path):
    knowledge, controls, retrieval = generate_rows()
    assert len(knowledge) == 180
    assert len(controls) == 60
    assert len(retrieval) == 360
    assert {domain: sum(row["domain"] == domain for row in knowledge) for domain in ("product", "process", "policy")} == {
        "product": 60, "process": 60, "policy": 60,
    }
    products = {}
    for split in ("dev", "validation", "test"):
        products[split] = {row["product_code"] for row in retrieval if row["split"] == split}
    assert not (products["dev"] & products["validation"])
    assert not (products["dev"] & products["test"])
    assert not (products["validation"] & products["test"])
    manifest = generate(tmp_path)
    assert manifest["files"]["knowledge.jsonl"]["rows"] == 180
    assert json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))["synthetic"] is True


@pytest.mark.asyncio
async def test_aliyun_client_contract_and_dimensions():
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["authorization"] == "Bearer test-only-key"
        calls.append(request.url.path)
        if request.url.path.endswith("/chat/completions"):
            return httpx.Response(200, json={
                "model": "qwen-plus",
                "choices": [{"message": {"content": "连接成功"}, "finish_reason": "stop"}],
                "usage": {"prompt_tokens": 2, "completion_tokens": 2},
            })
        body = json.loads(request.content)
        return httpx.Response(200, json={
            "model": "text-embedding-v4",
            "data": [{"index": index, "embedding": [float(index)] * 1024} for index, _ in enumerate(body["input"])],
            "usage": {"total_tokens": 3},
        })

    settings = Settings(
        _env_file=None,
        enable_paid_model_calls=True,
        dashscope_api_key=SecretStr("test-only-key"),
        dashscope_embedding_dimension=1024,
    )
    client = AliyunModelClient(settings, transport=httpx.MockTransport(handler))
    try:
        chat = await client.chat(ChatInput(messages=[ModelMessage(role="user", content="测试")]))
        vectors = await client.embed(["文档", "查询"])
    finally:
        await client.aclose()
    assert chat.text == "连接成功" and chat.input_tokens == 2
    assert len(vectors.vectors) == 2 and len(vectors.vectors[0]) == 1024
    assert calls == ["/compatible-mode/v1/chat/completions", "/compatible-mode/v1/embeddings"]


@pytest.mark.asyncio
async def test_aliyun_client_disabled_and_non_retryable_errors():
    disabled = AliyunModelClient(Settings(_env_file=None), transport=httpx.MockTransport(lambda _: httpx.Response(500)))
    with pytest.raises(PaidCallsDisabled):
        await disabled.embed(["test"])
    await disabled.aclose()

    count = 0
    def unauthorized(_: httpx.Request) -> httpx.Response:
        nonlocal count
        count += 1
        return httpx.Response(401, json={"error": {"message": "do-not-copy-body"}})

    settings = Settings(_env_file=None, enable_paid_model_calls=True, dashscope_api_key=SecretStr("secret"))
    client = AliyunModelClient(settings, transport=httpx.MockTransport(unauthorized))
    with pytest.raises(ModelServiceError) as caught:
        await client.embed(["test"])
    await client.aclose()
    assert caught.value.category == "authentication" and count == 1
    assert "do-not-copy-body" not in str(caught.value) and "secret" not in str(caught.value)


@pytest.mark.asyncio
async def test_admin_document_preview(runtime):
    user_headers = await login(runtime)
    assert (await runtime.client.post(
        "/api/v1/admin/documents/preview",
        headers=user_headers,
        files={"file": ("规则.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )).status_code == 403

    admin_headers = await login(runtime, "manager")
    response = await runtime.client.post(
        "/api/v1/admin/documents/preview",
        headers=admin_headers,
        files={"file": ("规则.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["source_type"] == "docx" and body["chunk_count"] >= 1
    assert body["notice"].startswith("仅完成解析预览")
